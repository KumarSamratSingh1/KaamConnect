import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

out_dir = os.path.abspath(r'KaamConnect_Project_Report')
img_dir = os.path.join(out_dir, 'screenshots')
docx_path = os.path.join(out_dir, 'KaamConnect_Project_Report.docx')
pdf_path = os.path.join(out_dir, 'KaamConnect_Project_Report.pdf')
md_path = os.path.join(out_dir, 'KaamConnect_Project_Report_Source.md')

doc = docx.Document()

# ----------------- PAGE SETUP (A4, 1-inch margins, 1.25 left for binding) -----------------
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.25)
section.right_margin = Inches(1.0)

# Colors
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_EMERALD = RGBColor(5, 150, 105)
COLOR_GRAY = RGBColor(71, 85, 105)

def set_cell_border(cell, **kwargs):
    """Set cell borders in Word table."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:left w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:bottom w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:right w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
        </w:tcBorders>
    '''.format(
        nsdecls('w'),
        kwargs.get('top', 'single'), kwargs.get('top_sz', '4'), kwargs.get('top_color', 'CCCCCC'),
        kwargs.get('left', 'single'), kwargs.get('left_sz', '4'), kwargs.get('left_color', 'CCCCCC'),
        kwargs.get('bottom', 'single'), kwargs.get('bottom_sz', '4'), kwargs.get('bottom_color', 'CCCCCC'),
        kwargs.get('right', 'single'), kwargs.get('right_sz', '4'), kwargs.get('right_color', 'CCCCCC'),
    ))
    tcPr.append(tcBorders)

def set_cell_background(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_title(text, level=0, space_before=0, space_after=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_chapter_heading(chap_num, chap_name):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run1 = p.add_run(f"CHAPTER {chap_num}\n")
    run1.font.name = 'Times New Roman'
    run1.font.size = Pt(16)
    run1.font.bold = True
    run1.font.color.rgb = COLOR_BLACK

    run2 = p.add_run(chap_name)
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(16)
    run2.font.bold = True
    run2.font.color.rgb = COLOR_BLACK
    return p

def add_h1(text, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_h2(text, space_before=10, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_p(text, line_spacing=1.5, space_after=6, justified=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justified else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_BLACK
    return p

def add_image_with_caption(img_path, caption_text, width_inches=5.8):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(10)
        run_cap.font.bold = True
        run_cap.font.color.rgb = COLOR_BLACK

print('Document framework initialized.')

