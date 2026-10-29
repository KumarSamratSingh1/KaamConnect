# -*- coding: utf-8 -*-
"""
Full Academic Project Report Generator for KaamConnect
Strictly adheres to university format specifications from Project Report Format.pdf:
- Page size: A4, 1.25 in left margin, 1.0 in top/bottom/right margin
- Font: Times New Roman throughout
- Chapter titles: 16pt Bold Centered
- Main headings: 14pt Bold Left Aligned
- Subheadings: 12pt Bold Left Aligned
- Body text: 12pt Justified, 1.5 line spacing
- Abstract & Certificate: 14pt, double line spacing
- Tables and figures numbered chapter-wise with 10pt Bold captions
- Full Prelims: Cover, Title, Certificate, Ack, TOC, List of Figures, List of Tables, Abstract, Graphical Abstract, Abbreviations
- Full Chapters 1 to 5 + References + Appendices (User Manual with screenshots, APIs, Schemas, Setup, Achievements)
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

out_dir = r'C:\Users\HP\OneDrive\5th sem\FSD II\Semester_Project\KaamConnect_Project_Report'
img_dir = os.path.join(out_dir, 'screenshots')
docx_path = os.path.join(out_dir, 'KaamConnect_Project_Report.docx')
pdf_path = os.path.join(out_dir, 'KaamConnect_Project_Report.pdf')
md_path = os.path.join(out_dir, 'KaamConnect_Project_Report_Source.md')

doc = docx.Document()

# ----------------- PAGE SETUP (A4, 1-inch margins, 1.25 left for binding) -----------------
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.25)
section.right_margin = Inches(1.0)

COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:left w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:bottom w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:right w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
        </w:tcBorders>
    '''.format(
        nsdecls('w'),
        kwargs.get('top', 'single'), kwargs.get('top_sz', '4'), kwargs.get('top_color', 'B0BEC5'),
        kwargs.get('left', 'single'), kwargs.get('left_sz', '4'), kwargs.get('left_color', 'B0BEC5'),
        kwargs.get('bottom', 'single'), kwargs.get('bottom_sz', '4'), kwargs.get('bottom_color', 'B0BEC5'),
        kwargs.get('right', 'single'), kwargs.get('right_sz', '4'), kwargs.get('right_color', 'B0BEC5'),
    ))
    tcPr.append(tcBorders)

def set_cell_background(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_chapter_heading(chap_num, chap_name):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.15
    run1 = p.add_run(f"CHAPTER {chap_num}\n")
    run1.font.name = 'Times New Roman'
    run1.font.size = Pt(16)
    run1.font.bold = True
    run1.font.color.rgb = COLOR_BLACK

    run2 = p.add_run(chap_name.upper())
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(16)
    run2.font.bold = True
    run2.font.color.rgb = COLOR_BLACK
    return p

def add_prelim_heading(title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(title.upper())
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_h1(text, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_h2(text, space_before=10, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_p(text, line_spacing=1.5, space_after=6, font_size=12, bold=False, italic=False, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = COLOR_BLACK
    return p

def add_bullet(text, bold_prefix='', line_spacing=1.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + ' ')
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(12)
    r_body.font.color.rgb = COLOR_BLACK
    return p

def add_image_with_caption(img_path, caption_text, width_inches=5.6):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(14)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(10)
        run_cap.font.bold = True
        run_cap.font.color.rgb = COLOR_BLACK

print('Step 1: Setup completed successfully.')

